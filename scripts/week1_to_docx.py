"""Convert a simple markdown file (week-1 deliverable) to a styled Word doc."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x1F, 0x38, 0x64)
ACCENT = RGBColor(0x2F, 0x54, 0x96)
CODE_INK = RGBColor(0x9C, 0x27, 0x4E)
HEADER_FILL = "1F3864"
BAND_FILL = "F4F6F9"
BODY_FONT = "Calibri"
HEAD_FONT = "Calibri Light"
MONO_FONT = "Consolas"
INLINE = re.compile(
    r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*(?!\*)[^*]+\*(?!\*))"
)


def set_cell_fill(cell, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(table, top=60, bottom=60, left=90, right=90) -> None:
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    table._tbl.tblPr.append(margins)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2F5496")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)


def write_inline(paragraph, text: str, *, base_size: Pt | None = None, bold_all: bool = False) -> None:
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.color.rgb = CODE_INK
            run.font.size = Pt(9)
        elif token.startswith("[") and "](" in token:
            label, _, url = token[1:-1].partition("](")
            add_hyperlink(paragraph, label, url)
            continue
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)
        if bold_all:
            run.bold = True
        if base_size is not None and run.font.size is None:
            run.font.size = base_size


def split_row(line: str) -> list[str]:
    body = line.strip().strip("|")
    return [p.strip() for p in re.split(r"(?<!\\)\|", body)]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    headers, body = rows[0], rows[1:]
    headless = all(not h for h in headers)
    data = body if headless else rows
    table = doc.add_table(rows=len(data), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    wide = len(headers) >= 3
    size = Pt(8.5) if wide else Pt(9.5)
    set_cell_margins(table)
    for r, row_cells in enumerate(data):
        is_header = not headless and r == 0
        for c in range(len(headers)):
            value = row_cells[c] if c < len(row_cells) else ""
            cell = table.cell(r, c)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            if is_header:
                write_inline(paragraph, value, base_size=size, bold_all=True)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_fill(cell, HEADER_FILL)
            else:
                write_inline(paragraph, value, base_size=size, bold_all=headless and c == 0)
                if r % 2 == (0 if not headless else 1):
                    set_cell_fill(cell, BAND_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color in (
        ("Heading 1", 18, INK),
        ("Heading 2", 13, ACCENT),
        ("Heading 3", 11, ACCENT),
    ):
        style = doc.styles[name]
        style.font.name = HEAD_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            write_inline(paragraph, heading.group(2))
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            if len(table_lines) >= 2:
                header = split_row(table_lines[0])
                rows = [header] + [split_row(line) for line in table_lines[2:]]
                add_table(doc, rows)
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            paragraph = doc.add_paragraph(style="List Bullet")
            write_inline(paragraph, bullet.group(1))
            i += 1
            continue
        paragraph = doc.add_paragraph()
        write_inline(paragraph, stripped)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"WROTE {out_path}")


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else root / "docs" / "week1-three-questions.md"
    target = Path(sys.argv[2]) if len(sys.argv) > 2 else root / "docs" / "Week1-Three-Questions.docx"
    convert(source, target)
